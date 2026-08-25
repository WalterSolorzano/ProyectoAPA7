import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2] / "python"))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from tests.test_cover_protection_roundtrip import _variant_table_first
from modules.scoped_apply import _cover_guard, _structural_cover
from docx import Document

raw = _variant_table_first()
import io
doc = Document(io.BytesIO(raw))
print("paras:", [p.text[:30] for p in doc.paragraphs])
print("structural:", len(_structural_cover(doc)))
try:
    from parsing.pre_classifier import pre_classify_elements
    from models import ElementModel, ElementType
    elems = [ElementModel(id=f"z{i}", type=ElementType.PARAGRAPH,
                          text=(p.text or "").strip(), original_text=(p.text or "").strip())
             for i, p in enumerate(doc.paragraphs)]
    cls = pre_classify_elements(elems)
    print("is_cover:", [bool(getattr(e, 'is_cover_section', False)) for e in cls])
except Exception as e:
    print("classifier EXC:", repr(e)[:200])
g = _cover_guard(doc)
print("guard:", g["protected"], g["detected"])

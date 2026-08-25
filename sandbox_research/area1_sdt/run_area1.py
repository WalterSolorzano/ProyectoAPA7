"""AREA 1 — Proteccion estructural de portada.

Experimentos:
  A) Envolver portada en w:sdt con lock y atacarla desde nuestro propio
     pipeline (lxml no respeta locks logicos: hipotesis a confirmar).
  B) Round-trip cero-diff: pipeline actual SIN scope de portada ->
     diff XML canonico de la zona. Debe dar identico.
  C) Trazado de deteccion de zona (pre_classify_elements) vs ground truth
     por variante del corpus.
"""
from __future__ import annotations

import io
import json
import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "python"))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree

from harness.docgen import OUT as CORPUS, BODY_SENTINEL, build_all
from harness.xmldiff import body_children_c14n, zone_diff

EVID = Path(__file__).resolve().parents[1] / "evidence"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ── A: SDT lock ───────────────────────────────────────────────────────────────

def wrap_cover_in_sdt(src: Path, dst: Path, cover_children: int,
                      lock_val: str = "sdtLocked") -> None:
    """Mueve los primeros N hijos de body dentro de un w:sdt bloqueado."""
    doc = docx.Document(str(src))
    body = doc.element.body
    children = list(body)[:cover_children]
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}>'
        f'<w:sdtPr><w:id w:val="777777"/><w:lock w:val="{lock_val}"/>'
        f'<w:docPartObj><w:docPartGallery w:val="WordAPA7-Portada"/></w:docPartObj></w:sdtPr>'
        f'<w:sdtContent/></w:sdt>'
    )
    content = sdt.find(qn("w:sdtContent"))
    body.insert(0, sdt)
    for child in children:
        content.append(child)  # append MUEVE el nodo (lxml)
    doc.save(str(dst))


def attack_naive_edit(path: Path) -> dict:
    """Intento de edicion directo con python-docx/lxml (nuestro propio poder)."""
    doc = docx.Document(str(path))
    changed = []
    for t in doc.element.body.iter(qn("w:t")):
        if t.text and "UNIVERSIDAD" in t.text:
            t.text = "TEXTO-VIOLADO-DESDE-LXML"
            changed.append("ok")
            break
    out = path.with_suffix(".attacked.docx")
    doc.save(str(out))
    # releer y verificar si el cambio persistio
    reread = docx.Document(str(out))
    persisted = any(
        (t.text or "") == "TEXTO-VIOLADO-DESDE-LXML"
        for t in reread.element.body.iter(qn("w:t"))
    )
    return {"naive_edit_attempted": bool(changed), "lock_stopped_us": not persisted}


def attack_pipeline(src_locked: Path, label: str) -> dict:
    """scoped_apply con scopes que NO incluyen portada (no existe tal scope)."""
    from modules.scoped_apply import apply_scopes
    raw = src_locked.read_bytes()
    out_bytes, summary = apply_scopes(raw, ["texto", "tablas_imagenes"], {})
    out_path = src_locked.with_suffix(f".scoped_{label}.docx")
    out_path.write_bytes(out_bytes)
    before = body_children_c14n(str(src_locked))
    after = body_children_c14n(str(out_path))
    return {"pipeline_touched_sdt_zone": not _sdt_zone_equal(before, after)}


def _sdt_zone_equal(before: list[str], after: list[str]) -> bool:
    def sdt_zone(chunks):
        for c in chunks:
            if "<w:sdt" in c:
                return c
        return None
    return sdt_zone(before) == sdt_zone(after)


# ── B: round-trip cero diff (pipeline actual, sin SDT) ────────────────────────

def roundtrip_zero_diff(variant: str, meta: dict) -> dict:
    from modules.scoped_apply import apply_scopes
    src = CORPUS / f"{variant}.docx"
    raw = src.read_bytes()
    t0 = time.perf_counter()
    out_bytes, summary = apply_scopes(raw, ["texto", "tablas_imagenes"], {})
    ms = (time.perf_counter() - t0) * 1000
    out = src.with_suffix(".rt.docx")
    out.write_bytes(out_bytes)
    rep = zone_diff(body_children_c14n(str(src)), body_children_c14n(str(out)),
                    meta["cover_children"])
    return {"variant": variant, "ms": round(ms), **rep, "summary": summary}


# ── C: trazado deteccion de zonas vs ground truth ─────────────────────────────

def ground_truth_paragraphs(path: Path) -> tuple[list[str], int]:
    """Parrafos top-level (como los enviaria Word) + cuantos son de portada."""
    from parsing.docx_parser import _extract_paragraph_text_with_footnotes
    with __import__("zipfile").ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    paras, cover_n, seen_sentinel = [], 0, False
    for child in body:
        tag = etree.QName(child).localname
        if tag == "tbl":
            continue  # Office.js las cuenta aparte; registro aparte
        if tag != "p":
            continue
        text, _ = _extract_paragraph_text_with_footnotes(child)
        paras.append(text)
        if not seen_sentinel:
            cover_n += 1
            if BODY_SENTINEL in text:
                seen_sentinel = True
    return paras, cover_n


def trace_zones(variant: str) -> dict:
    from parsing.pre_classifier import pre_classify_elements
    from models import ElementModel, ElementType
    src = CORPUS / f"{variant}.docx"
    paras, gt_cover = ground_truth_paragraphs(src)
    elems = [ElementModel(id=f"z{i}", type=ElementType.PARAGRAPH, text=(t or "").strip(),
                          original_text=(t or "").strip()) for i, t in enumerate(paras)]
    classified = pre_classify_elements(elems)
    is_cover = [bool(getattr(e, "is_cover_section", False)) for e in classified]
    detected = sum(is_cover[:gt_cover])
    false_pos = sum(is_cover[gt_cover:])
    return {
        "variant": variant,
        "ground_truth_cover_paras": gt_cover,
        "detected_inside_gt": detected,
        "false_positives_after": false_pos,
        "misses": gt_cover - detected,
        "body_start_idx": next((i for i, c in enumerate(is_cover) if not c), len(is_cover)),
        "exact": detected == gt_cover and false_pos == 0,
    }


# ── main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    build_all()
    results: dict = {"A_sdt_lock": [], "B_roundtrip": [], "C_zones": []}

    # B primero: baseline sin SDT
    meta = json.loads((CORPUS / "meta.json").read_text(encoding="utf-8"))
    for variant, m in meta.items():
        try:
            results["B_roundtrip"].append(roundtrip_zero_diff(variant, m))
        except Exception as e:
            results["B_roundtrip"].append({"variant": variant, "error": repr(e)[:200]})

    # C: deteccion de zonas
    for variant in meta:
        try:
            results["C_zones"].append(trace_zones(variant))
        except Exception as e:
            results["C_zones"].append({"variant": variant, "error": repr(e)[:200]})

    # A: SDT lock sobre plaintext y mixed
    for variant in ("plaintext", "mixed"):
        src = CORPUS / f"{variant}.docx"
        for lock in ("sdtLocked", "contentLocked"):
            locked = CORPUS / f"{variant}_locked_{lock}.docx"
            wrap_cover_in_sdt(src, locked, meta[variant]["cover_children"], lock)
            res = {"variant": variant, "lock": lock}
            res.update(attack_naive_edit(locked))
            try:
                res.update(attack_pipeline(locked, lock))
            except Exception as e:
                res["pipeline_error"] = repr(e)[:200]
            results["A_sdt_lock"].append(res)

    EVID.mkdir(exist_ok=True)
    (EVID / "area1.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

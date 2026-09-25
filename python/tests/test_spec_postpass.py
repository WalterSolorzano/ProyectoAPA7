"""Tests del post-paso: estilos heading, tarjetas de anexo, bordes por caption."""
import base64
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import docx
import pytest
from docx.shared import Pt
from preset_store import HeadingLevelPreset
from spec_postpass import (
    EquipmentCardItem,
    append_equipment_cards,
    apply_heading_styles,
    apply_table_border_override,
)

# PNG valido 1x1 (transparente)
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYGBg"
    "AAAABQABh6FO1AAAAABJRU5ErkJggg==")


@pytest.fixture
def png_path(tmp_path):
    p = tmp_path / "eq.png"
    p.write_bytes(_PNG)
    return p


def _new_doc(tmp_path):
    d = docx.Document()
    d.add_paragraph("hola")
    path = tmp_path / "out.docx"
    d.save(str(path))
    return path


def test_apply_heading_styles_sets_page_break(tmp_path):
    path = _new_doc(tmp_path)
    apply_heading_styles(path, {"1": HeadingLevelPreset(
        page_break_before=True, size_pt=14, bold=True,
        font_family="Times New Roman")})
    d = docx.Document(str(path))
    st = d.styles["Heading 1"]
    assert st.paragraph_format.page_break_before is True
    assert st.font.size == Pt(14)
    assert st.font.name == "Times New Roman"
    # Tema limpio (F-10): sin w:asciiTheme en rFonts
    from docx.oxml.ns import qn
    rPr = st.element.find(qn("w:rPr"))
    rFonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
    assert rFonts is None or rFonts.get(qn("w:asciiTheme")) is None


def test_append_equipment_cards(tmp_path, png_path):
    path = _new_doc(tmp_path)
    append_equipment_cards(path, [EquipmentCardItem(
        number="A1", title="Refrigerador", image=str(png_path),
        specs={"Potencia": "120 W"})])
    d = docx.Document(str(path))
    texts = [p.text for p in d.paragraphs]
    assert any("Figura A1" in t for t in texts)
    assert len(d.tables) >= 1


def test_table_border_override_found(tmp_path):
    path = _new_doc(tmp_path)
    d = docx.Document(str(path))
    d.add_paragraph("Tabla 1").runs[0].bold = True
    d.add_table(rows=2, cols=2)
    d.save(str(path))
    assert apply_table_border_override(path, "Tabla 1", "grid") is True


def test_table_border_override_missing_returns_false(tmp_path):
    path = _new_doc(tmp_path)
    assert apply_table_border_override(path, "Tabla 99", "grid") is False

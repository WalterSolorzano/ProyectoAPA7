"""Test F3: no duplicar la sección de Referencias en la ruta de rebuild.

Cuando el DocumentModel incluye un heading "Referencias"/"Bibliografía" más
las entradas de referencia (lo que el parser extrae del original) Y además se
provee la lista `references`, el generador NO debe escribir la sección dos
veces. Debe aparecer un único heading "Referencias" (escrito por
format_apa_referencias_section al final).

Cubre:
* Ruta in-place (generator.py sobre original.docx): el heading + entradas del
  modelo se eliminan del cuerpo y se reescriben una sola vez desde la lista.
* Ruta layered (layered_generator.py, WORDAPA7_LAYERED_GEN=1): el heading +
  entradas del modelo se omiten en PHASE 3 y se escriben una sola vez en
  PHASE 4.
* Sin lista `references`: no hay dedup (no se borra el heading del modelo).
"""
import io
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import docx

from generation.generator import generate_apa7_docx
from models import (
    APARuleSet,
    DocumentModel,
    ElementModel,
    ElementType,
    ReferenciaModel,
)

REF_GARCIA = (
    "García, M. (2023). Inteligencia artificial en el aula: Un estudio "
    "comparativo. Revista de Educación Superior, 45(2), 123-145."
)
REF_LOPEZ = (
    "López, R. (2021). Modelos de lenguaje y procesamiento de texto "
    "académico. Editorial Universitaria."
)


def _references_list():
    return [
        ReferenciaModel(
            id="r1", authors=["García", "M."], year="2023",
            title="Inteligencia artificial en el aula",
            source="Revista de Educación Superior, 45(2), 123-145",
            raw_text=REF_GARCIA,
        ),
        ReferenciaModel(
            id="r2", authors=["López", "R."], year="2021",
            title="Modelos de lenguaje y procesamiento de texto académico",
            source="Editorial Universitaria",
            raw_text=REF_LOPEZ,
        ),
    ]


def _model_with_refs_section():
    """Modelo con cuerpo + heading 'Referencias' + entradas de referencia."""
    return [
        ElementModel(id="b1", type=ElementType.PARAGRAPH, text="Cuerpo del documento."),
        ElementModel(id="h1", type=ElementType.HEADING, heading_level=1, text="Referencias"),
        ElementModel(
            id="r1", type=ElementType.PARAGRAPH, text=REF_GARCIA,
            pre_classifier_rule="reference_item",
        ),
        ElementModel(
            id="r2", type=ElementType.PARAGRAPH, text=REF_LOPEZ,
            pre_classifier_rule="reference_item",
        ),
    ]


def _count_referencias(doc):
    return sum(1 for p in doc.paragraphs if p.text.strip() == "Referencias")


# ── Ruta in-place (original.docx) ────────────────────────────────────────────


def test_single_references_heading_when_list_provided_inplace(tmp_path):
    """Con original.docx + lista `references`: un solo heading 'Referencias'."""
    # original.docx con cuerpo + Referencias + entradas (orden = orden del modelo)
    orig = docx.Document()
    orig.add_paragraph("Cuerpo del documento.")
    orig.add_paragraph("Referencias")
    orig.add_paragraph(REF_GARCIA)
    orig.add_paragraph(REF_LOPEZ)
    (tmp_path / "original.docx").write_bytes(_docx_bytes(orig))

    doc = DocumentModel(
        session_id="s1", file_name="t.docx",
        elements=_model_with_refs_section(),
    )
    out = tmp_path / "out.docx"
    generate_apa7_docx(doc, out, rules=APARuleSet(), references=_references_list())

    d = docx.Document(str(out))
    assert _count_referencias(d) == 1, (
        "Debe haber exactamente un heading 'Referencias' en la salida "
        f"(hubo {_count_referencias(d)})"
    )
    all_text = "\n".join(p.text for p in d.paragraphs)
    assert "García" in all_text, "La entrada de García debe aparecer en la salida"
    assert "López" in all_text, "La entrada de López debe aparecer en la salida"


def test_references_not_deduped_when_no_list_inplace(tmp_path):
    """Sin lista `references` no se elimina el heading del modelo (no double,
    no pérdida): sigue habiendo un único 'Referencias' formateado in-place."""
    orig = docx.Document()
    orig.add_paragraph("Cuerpo del documento.")
    orig.add_paragraph("Referencias")
    orig.add_paragraph(REF_GARCIA)
    orig.add_paragraph(REF_LOPEZ)
    (tmp_path / "original.docx").write_bytes(_docx_bytes(orig))

    doc = DocumentModel(
        session_id="s2", file_name="t.docx",
        elements=_model_with_refs_section(),
    )
    out = tmp_path / "out.docx"
    # Sin references: la sección del modelo se formatea in-place tal cual.
    generate_apa7_docx(doc, out, rules=APARuleSet())

    d = docx.Document(str(out))
    assert _count_referencias(d) == 1, (
        "Sin lista de referencias debe haber un único 'Referencias' (el del modelo)"
    )
    all_text = "\n".join(p.text for p in d.paragraphs)
    assert "García" in all_text and "López" in all_text


# ── Ruta layered (from-scratch) ───────────────────────────────────────────────


def test_single_references_heading_when_list_provided_layered(tmp_path, monkeypatch):
    """Ruta layered (WORDAPA7_LAYERED_GEN=1): un solo heading 'Referencias'."""
    monkeypatch.setenv("WORDAPA7_LAYERED_GEN", "1")

    doc = DocumentModel(
        session_id="s3", file_name="t.docx",
        elements=_model_with_refs_section(),
    )
    out = tmp_path / "out.docx"
    generate_apa7_docx(doc, out, rules=APARuleSet(), references=_references_list())

    d = docx.Document(str(out))
    assert _count_referencias(d) == 1, (
        "Ruta layered: debe haber exactamente un heading 'Referencias' "
        f"(hubo {_count_referencias(d)})"
    )
    all_text = "\n".join(p.text for p in d.paragraphs)
    assert "García" in all_text and "López" in all_text


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

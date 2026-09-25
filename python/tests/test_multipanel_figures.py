"""Test para figuras multipanel APA 7 (a, b)"""
import pathlib
import sys

import docx
from PIL import Image

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from generation.generator import generate_apa7_docx
from models import APARuleSet, DocumentModel, ElementModel, ElementType, ImageModel, SubfigureModel


def _make_test_png(path, color="blue"):
    Image.new("RGB", (100, 100), color).save(str(path), format="PNG")
    return str(path)

def test_multipanel_figure_generation(tmp_path):
    img_a = _make_test_png(tmp_path / "panel_a.png", "red")
    img_b = _make_test_png(tmp_path / "panel_b.png", "green")

    sub_a = SubfigureModel(
        id="sub1", label="(a)", title="Vista general",
        relative_url="/api/images/test/panel_a.png",
        file_path=img_a, filename="panel_a.png"
    )
    sub_b = SubfigureModel(
        id="sub2", label="(b)", title="Detalle técnico",
        relative_url="/api/images/test/panel_b.png",
        file_path=img_b, filename="panel_b.png"
    )

    image_info = ImageModel(
        element_id="img_multi_1",
        file_path=img_a,
        filename="panel_a.png",
        caption="Comparación de prototipos",
        note="Datos obtenidos en laboratorio.",
        figure_number=1,
        design_style="multipanel",
        subfigures=[sub_a, sub_b]
    )

    doc_model = DocumentModel(
        session_id="multi_sess",
        file_name="multipanel.docx",
        elements=[
            ElementModel(id="img_multi_1", type=ElementType.IMAGE, text="", image_info=image_info)
        ]
    )

    out_docx = tmp_path / "multipanel_out.docx"
    generate_apa7_docx(doc_model, out_docx, rules=APARuleSet())

    assert out_docx.exists()
    d = docx.Document(str(out_docx))

    # Debe contener una tabla de 2 columnas para el multipanel
    tables = d.tables
    assert len(tables) >= 1
    tbl = tables[0]
    assert len(tbl.columns) == 2
    # La fila 1 debe contener los subtítulos '(a) Vista general' y '(b) Detalle técnico'
    cell_a = tbl.cell(1, 0).text
    cell_b = tbl.cell(1, 1).text
    assert "(a)" in cell_a and "Vista general" in cell_a
    assert "(b)" in cell_b and "Detalle técnico" in cell_b

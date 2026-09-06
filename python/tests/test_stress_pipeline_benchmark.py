"""
Stress Pipeline Benchmark Test:
Verifica que el motor WordAPA7 procese, parsee y transforme in-place
todos los escenarios complejos de DocxStressLab sin errores ni corrupción de archivo.
"""

import os
from pathlib import Path
import pytest
from docx import Document
from tools.stress_doc_generator import generate_all_stress_docs
from parsing.docx_parser import parse_docx_bytes
from models import APARuleSet
from generation.inplace_editor import apply_inplace
from modules.doc_auditor import audit_document_heuristic

@pytest.fixture(scope="session")
def stress_docs():
    return generate_all_stress_docs()

def test_stress_pipeline_all_scenarios(stress_docs, tmp_path):
    rules = APARuleSet()
    for scenario_name, doc_path in stress_docs.items():
        assert os.path.exists(doc_path), f"El archivo {doc_path} no existe"

        with open(doc_path, "rb") as f:
            bytes_data = f.read()

        # 1. Parseo
        session_id = f"stress_{scenario_name}"
        model = parse_docx_bytes(bytes_data, f"{scenario_name}.docx", session_id=session_id, storage_dir=tmp_path)
        assert model is not None
        assert len(model.elements) > 0

        # 2. Auditoría heurística
        audit = audit_document_heuristic(model)
        assert audit is not None

        # 3. Transformación in-place a APA 7
        out_docx_path = tmp_path / f"output_{scenario_name}.docx"
        res_path = apply_inplace(
            original_path=Path(doc_path),
            out_path=out_docx_path,
            doc_model=model,
            rules=rules,
            scopes=["texto", "tablas_imagenes", "bibliografia"]
        )

        assert os.path.exists(res_path)
        # 4. Validar que el archivo resultante es un DOCX válido
        out_doc = Document(res_path)
        assert len(out_doc.paragraphs) > 0 or len(out_doc.tables) > 0

"""Regresión FASE 1.2 — portada intacta tras scopes que NO la piden.

Evidencia origen: docs/evaluacion-tecnologica/EVALUACION_TECNOLOGICA.md S4
(area1.json B_roundtrip: zone_identical=false en 5/5 variantes).
Vectores cubiertos: V2 sangría en párrafos de portada, V3 reformateo de
tablas/captions dentro de la zona.
"""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import io
import zipfile

from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SENTINEL = "INICIO-DEL-CUERPO-SENTINEL"


def _c14n_children(path_bytes: bytes) -> list[str]:
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(path_bytes)).read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    return [etree.canonicalize(etree.tostring(c).decode()) for c in body]


def _finish(doc: Document):
    doc.add_paragraph(SENTINEL)
    doc.add_heading("Introduccion", level=1)
    doc.add_paragraph("Cuerpo del trabajo con texto suficiente.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _variant_plaintext() -> bytes:
    d = Document()
    for line in ("UNIVERSIDAD NACIONAL", "Facultad de Ingenieria", "Autor Uno"):
        p = d.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    out = _finish(d)
    # ground truth: 3 hijos de portada
    return out


def _variant_table_first() -> bytes:
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "LOGO"
    t.cell(0, 1).text = "UNIVERSIDAD CON TABLA"
    p = d.add_paragraph("Titulo sobre tabla")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    out = _finish(d)
    return out


def _variant_textbox() -> bytes:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    d = Document()
    host = d.add_paragraph()
    NS = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    )
    ac = parse_xml(
        f'<mc:AlternateContent {NS}>'
        "<mc:Choice Requires=\"wps\"><w:drawing><wp:anchor distT=\"0\" distB=\"0\" distL=\"0\" distR=\"0\" "
        "simplePos=\"0\" relativeHeight=\"1\" behindDoc=\"0\" locked=\"0\" layoutInCell=\"1\" allowOverlap=\"1\">"
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="3600000" cy="900000"/><wp:wrapNone/>'
        '<wp:docPr id="91" name="TxBoxPortada"/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>AUTOR-TEXTBOX</w:t></w:r></w:p></w:txbxContent>"
        "</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:shape xmlns:v=\"urn:schemas-microsoft-com:vml\" style=\"width:283pt;height:70pt\">"
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>AUTOR-TEXTBOX</w:t></w:r></w:p></w:txbxContent></v:textbox>'
        "</v:shape></w:pict></mc:Fallback></mc:AlternateContent>"
    )
    host._p.append(ac)
    t = d.add_paragraph("TITULO-TRAS-TEXTBOX")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return _finish(d)


def _zone(before: bytes, after: bytes, cover_children: int) -> bool:
    b, a = _c14n_children(before)[:cover_children], _c14n_children(after)[:cover_children]
    return b == a


def test_scope_texto_y_tablas_dejan_portada_intacta():
    from modules.scoped_apply import apply_scopes
    casos = [(_variant_plaintext(), 3), (_variant_table_first(), 2), (_variant_textbox(), 2)]
    for raw, cover_n in casos:
        out, summary = apply_scopes(raw, ["texto", "tablas_imagenes"], {})
        assert _zone(raw, out, cover_n), (
            f"portada modificada (cover_children={cover_n}); summary={summary}"
        )
        assert summary.get("cover_guard", {}).get("protected", 0) >= 1


def test_sin_deteccion_comporta_como_antes_sin_inventar():
    """Documento sin senales de portada: el guard no protege nada y no rompe."""
    from modules.scoped_apply import apply_scopes
    d = Document()
    d.add_paragraph("Parrafo uno normal")
    d.add_paragraph("Parrafo dos normal")
    buf = io.BytesIO(); d.save(buf); raw = buf.getvalue()
    out, summary = apply_scopes(raw, ["texto"], {})
    assert summary["cover_guard"]["protected"] == 0
    assert len(_c14n_children(out)) == len(_c14n_children(raw))

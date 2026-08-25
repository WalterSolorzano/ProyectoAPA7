"""Regresión FASE 1.1 — colisión de numId en bullet_engine.

Evidencia origen: docs/evaluacion-tecnologica/EVALUACION_TECNOLOGICA.md S2
(PoC area4.json misformat_confirmed=true: viñeta hereda decimal del usuario).
"""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _make_doc_with_user_decimal_list() -> Document:
    """Doc con numbering.xml propio: numId=1 -> abstract DECIMAL (lista del usuario)."""
    from generation.bullet_engine import install_numbering_part
    base = Document()  # plantilla default (trae numbering 1..9); se REEMPLAZA abajo
    d = install_numbering_part(
        base,
        [
            {
                "abstract_num_id": "9",
                "num_id": "1",
                "fmt": "decimal",
                "lvl_text": "%1.",
            }
        ],
    )
    p1 = d.add_paragraph("Usuario item A")
    p2 = d.add_paragraph("Usuario item B")
    for p in (p1, p2):
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId"); nid.set(qn("w:val"), "1")
        numPr.append(ilvl); numPr.append(nid); pPr.append(numPr)
    return d


def _resolved_fmt_for(doc: Document, paragraph) -> str | None:
    """Formato del abstractNum al que resuelve el numId del parrafo."""
    npart = None
    for rid, rel in doc.part.rels.items():
        if rel.reltype.endswith("/numbering"):
            npart = rel.target_part
            break
    if npart is None:
        return None
    num_el = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
    assert num_el is not None, "parrafo sin numPr"
    used = num_el.find(qn("w:numId")).get(qn("w:val"))
    numbering = npart.element
    aid = None
    for n in numbering.findall(qn("w:num")):
        if n.get(qn("w:numId")) == used:
            aid = n.find(qn("w:abstractNumId")).get(qn("w:val"))
    if aid is None:
        return None
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == aid:
            fmt = a.find(f"{qn('w:lvl')}/{qn('w:numFmt')}")
            return fmt.get(qn("w:val")) if fmt is not None else None
    return None


def test_bullet_does_not_hijack_user_decimal_list():
    from generation.bullet_engine import format_bullet_item
    d = _make_doc_with_user_decimal_list()
    nuevo = d.add_paragraph("Vineta APA nueva")
    format_bullet_item(nuevo, "Texto de la vineta", level=1)
    fmt = _resolved_fmt_for(d, nuevo)
    assert fmt == "bullet", f"la vineta heredo {fmt!r} del usuario (colision numId)"
    # los items del usuario siguen en su lista decimal intacta
    assert _resolved_fmt_for(d, d.paragraphs[0]) == "decimal"
    assert _resolved_fmt_for(d, d.paragraphs[1]) == "decimal"


def test_consecutive_bullets_reuse_same_new_numid():
    from generation.bullet_engine import format_bullet_item
    d = _make_doc_with_user_decimal_list()
    a = d.add_paragraph("Uno"); format_bullet_item(a, "Uno", level=1)
    b = d.add_paragraph("Dos"); format_bullet_item(b, "Dos", level=1)
    fa, fb = _resolved_fmt_for(d, a), _resolved_fmt_for(d, b)
    assert fa == fb == "bullet"
    ia = a._p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")).get(qn("w:val"))
    ib = b._p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")).get(qn("w:val"))
    assert ia == ib, "cada vineta creo un numId distinto (debe reutilizar)"


def test_numbered_item_gets_own_decimal_definition():
    from generation.bullet_engine import format_numbered_item
    d = _make_doc_with_user_decimal_list()
    n = d.add_paragraph("Paso nuevo")
    format_numbered_item(n, "Paso nuevo", level=1)
    # no debe compartir el numId=1 del usuario (reiniciaria/continuaria su lista)
    used = n._p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")).get(qn("w:val"))
    assert used != "1", "lista numerada nueva comparte numId del usuario"

"""Test: Las imágenes tienen keep_together/widow_control (D3) y su altura se
limita al alto utilizable de la página (D4).

D3 — keep_together/widow_control en párrafos de imagen:
    * Ruta "nueva imagen" (format_apa_figure): la imagen insertada lleva
      keep_together=True y widow_control=True.
    * Ruta "in-place" (_apply_image_design_style): una imagen existente en el
      original.docx también recibe ambos flags.

D4 — tope de altura:
    * Una imagen más alta que el área imprimible se escala proporcionalmente
      para que su altura (wp:extent cy) no exceda el alto utilizable de la
      página (page_height − top_margin − bottom_margin).
"""
import io
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import docx
from docx.oxml.ns import qn
from docx.shared import Inches

from PIL import Image

from generation.generator import generate_apa7_docx
from models import (
    APARuleSet,
    DocumentModel,
    ElementModel,
    ElementType,
    ImageModel,
)

BLIP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def _make_png(path, width=120, height=120, color="white"):
    """Crea una PNG válida en disco (requerida por python-docx add_picture)."""
    Image.new("RGB", (width, height), color).save(str(path), format="PNG")
    return path


def _image_paragraphs(doc):
    """Lista de párrafos que contienen un blip (imagen real)."""
    out = []
    for p in doc.paragraphs:
        if p._element.findall(f".//{BLIP_NS}"):
            out.append(p)
    return out


def _extent_cy(p):
    """Altura (EMU) del primer wp:extent del drawing del párrafo, o None."""
    for drawing in p._element.iter(qn("w:drawing")):
        for ext in drawing.iter(qn("wp:extent")):
            try:
                return int(ext.get("cy", "0"))
            except (ValueError, TypeError):
                return None
    return None


# ── D3: keep_together / widow_control (ruta nueva imagen) ───────────────────


def test_image_paragraph_has_keep_together(tmp_path):
    """La imagen insertada (sin original.docx) lleva keep_together/widow_control."""
    img_path = _make_png(tmp_path / "fig.png")
    img = ImageModel(
        element_id="i1", file_path=str(img_path), filename="fig.png",
        width_cm=12, height_cm=8, design_style="standard",
    )
    doc = DocumentModel(
        session_id="s1", file_name="t.docx",
        elements=[ElementModel(id="i1", type=ElementType.IMAGE, text="", image_info=img)],
    )
    out = tmp_path / "out.docx"
    generate_apa7_docx(doc, out, rules=APARuleSet())

    d = docx.Document(str(out))
    img_paras = _image_paragraphs(d)
    assert img_paras, "Should have at least one paragraph with an image"
    for p in img_paras:
        assert p.paragraph_format.keep_together is True, "Image should have keep_together=True"
        assert p.paragraph_format.widow_control is True, "Image should have widow_control=True"


# ── D3 + D4: ruta in-place (original.docx con una imagen muy alta) ───────────


def test_inplace_image_keep_together_and_height_clamped(tmp_path):
    """Una imagen existente más alta que la página se limita al alto utilizable
    y recibe keep_together/widow_control (in-place sobre original.docx)."""
    # Imagen alta (proporción vertical) insertada a 20" de alto en el original.
    img_path = _make_png(tmp_path / "tall.png", width=200, height=800, color="white")
    orig = docx.Document()
    run = orig.add_paragraph().add_run()
    run.add_picture(str(img_path), height=Inches(20))
    (tmp_path / "original.docx").write_bytes(_docx_bytes(orig))

    img = ImageModel(
        element_id="i1", file_path=str(img_path), filename="tall.png",
        width_cm=8, height_cm=40, design_style="standard",
    )
    doc = DocumentModel(
        session_id="s2", file_name="t.docx",
        elements=[ElementModel(id="i1", type=ElementType.IMAGE, text="", image_info=img)],
    )
    out = tmp_path / "out.docx"
    generate_apa7_docx(doc, out, rules=APARuleSet())

    d = docx.Document(str(out))
    img_paras = _image_paragraphs(d)
    assert img_paras, "Should have at least one paragraph with an image"

    sec = d.sections[0]
    usable_emu = int(sec.page_height) - int(sec.top_margin) - int(sec.bottom_margin)

    for p in img_paras:
        # D3: la imagen no se parte entre páginas.
        assert p.paragraph_format.keep_together is True, "Image should have keep_together=True"
        assert p.paragraph_format.widow_control is True, "Image should have widow_control=True"

        # D4: la altura no excede el área imprimible (tolerancia de redondeo).
        cy = _extent_cy(p)
        assert cy is not None, "Image paragraph should expose a wp:extent"
        assert cy <= usable_emu + 200, (
            f"Image height {cy} EMU exceeds usable height {usable_emu} EMU (should be clamped)"
        )
        # Confirmar que realmente se redujo (el original era 20" = 18_288_000 EMU).
        assert cy < 18_000_000, f"Image height {cy} was not scaled down from 20 inches"


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

"""
Pruebas unitarias para el Motor de Trasplante Quirúrgico de Portada OpenXML (Metodología A + B)
"""

import pytest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

from generation.openxml_cover import (
    find_all_rel_ids,
    create_isolated_section_break,
    transfer_cover_relationships,
    splice_cover_with_openxml,
)


def test_create_isolated_section_break():
    sect_br = create_isolated_section_break()
    assert sect_br is not None
    assert sect_br.tag.endswith("p")
    sect_types = sect_br.xpath("w:pPr/w:sectPr/w:type")
    assert len(sect_types) > 0
    assert sect_types[0].attrib[qn("w:val")] == "nextPage"


def test_splice_cover_with_openxml_real_docs(tmp_path):
    # Crear un docx original con párrafos de portada
    orig_path = tmp_path / "original.docx"
    gen_path = tmp_path / "generated.docx"
    out_path = tmp_path / "output.docx"

    doc_orig = Document()
    doc_orig.add_heading("UNIVERSIDAD NACIONAL DE INGENIERÍA", level=1)
    doc_orig.add_paragraph("FACULTAD DE ELECTROTECNIA Y COMPUTACIÓN")
    doc_orig.add_paragraph("Docente: Ing. Juan Carlos Aburto Poveda")
    doc_orig.save(str(orig_path))

    doc_gen = Document()
    doc_gen.add_heading("1. Introducción", level=1)
    doc_gen.add_paragraph("Este es el cuerpo del trabajo en APA 7.")
    doc_gen.save(str(gen_path))

    success = splice_cover_with_openxml(orig_path, gen_path, out_path, body_start_idx=3)
    assert success == True
    assert out_path.exists()

    result_doc = Document(str(out_path))
    texts = [p.text for p in result_doc.paragraphs if p.text.strip()]
    assert any("UNIVERSIDAD" in t for t in texts)
    assert any("Docente" in t for t in texts)
    assert any("Introducción" in t for t in texts)

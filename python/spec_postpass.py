"""Post-paso sobre el docx generado por POST /api/spec.

Aplica: estilos nativos Heading N desde preset (pagina nueva, tipografia,
limpieza de atributos de tema F-10), bordes de tabla por etiqueta de caption,
y tarjetas de anexo (equipment cards) al final del documento.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

from generation.image_handler import add_apa_equipment_card
from generation.table_engine import set_table_borders
from preset_store import HeadingLevelPreset
from spec_dsl import EquipmentCardItem  # re-export para consumidores

__all__ = ["EquipmentCardItem", "apply_heading_styles",
           "append_equipment_cards", "apply_table_border_override"]


def _strip_theme_and_set(style, font_name: str) -> None:
    """Fija w:ascii/w:hAnsi directos y elimina w:*Theme (fallo F-10)."""
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)


def apply_heading_styles(doc_path: Path, levels: dict) -> None:
    """Aplica preset de headings a los estilos nativos Heading 1..5."""
    if not levels:
        return
    d = docx.Document(str(doc_path))
    changed = False
    for lvl, p in levels.items():
        try:
            style = d.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        if p.font_family:
            style.font.name = p.font_family
            _strip_theme_and_set(style, p.font_family)
        if p.size_pt is not None:
            style.font.size = Pt(p.size_pt)
        if p.bold is not None:
            style.font.bold = p.bold
        if p.italic is not None:
            style.font.italic = p.italic
        if p.page_break_before is not None:
            style.paragraph_format.page_break_before = p.page_break_before
        if p.keep_with_next is not None:
            style.paragraph_format.keep_with_next = p.keep_with_next
        changed = True
    if changed:
        d.save(str(doc_path))


def append_equipment_cards(doc_path: Path, cards: list, rules=None) -> None:
    """Anexa tarjetas de equipo (foto + ficha nativa) al final del documento."""
    if not cards:
        return
    d = docx.Document(str(doc_path))
    for c in cards:
        add_apa_equipment_card(d, c.number, c.title, c.image, c.specs, rules)
    d.save(str(doc_path))


def apply_table_border_override(doc_path: Path, caption_label: str,
                                border_style: str) -> bool:
    """Busca la tabla cuyo parrafo previo (<=3) es la etiqueta y cambia bordes."""
    d = docx.Document(str(doc_path))
    for table in d.tables:
        prev = table._element.getprevious()
        hops = 0
        while prev is not None and hops < 3:
            if prev.tag == qn("w:p"):
                text = "".join(node.text or "" for node in prev.iter(qn("w:t")))
                if text.strip() == caption_label:
                    set_table_borders(table, border_style)
                    d.save(str(doc_path))
                    return True
            prev = prev.getprevious()
            hops += 1
    return False
